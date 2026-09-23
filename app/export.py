"""
export.py
---------
Export the tracked-documents report to CSV (always available, stdlib only)
or to a formatted Excel workbook if the optional 'openpyxl' package is
installed (pip install openpyxl).
"""

import csv

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False

COLUMNS = ["category", "company_name", "owner_name", "document", "number",
           "expiry", "status", "days_left"]
HEADERS = ["Type", "Company", "Employee/Company Name", "Document", "Document No.",
           "Expiry Date", "Status", "Days Left"]

STATUS_FILL = {
    "Expired": "F1948A",
    "Critical": "F5B7B1",
    "Warning": "FAD7A0",
    "Upcoming": "F9E79F",
    "Valid": "ABEBC6",
    "No Date": "D5D8DC",
}


def export_csv(items, filepath):
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
        for it in items:
            writer.writerow([it.get(c, "") for c in COLUMNS])
    return filepath


def export_excel(items, filepath):
    if not HAVE_OPENPYXL:
        raise RuntimeError("openpyxl is not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Document Expiry Report"

    ws.append(HEADERS)
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for it in items:
        row = [it.get(c, "") for c in COLUMNS]
        ws.append(row)
        r = ws.max_row
        fill_color = STATUS_FILL.get(it.get("status"), "FFFFFF")
        for col_idx in range(1, len(HEADERS) + 1):
            ws.cell(row=r, column=col_idx).fill = PatternFill(
                start_color=fill_color, end_color=fill_color, fill_type="solid")

    widths = [12, 26, 24, 28, 18, 14, 12, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(filepath)
    return filepath
